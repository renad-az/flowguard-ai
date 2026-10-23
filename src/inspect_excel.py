import argparse
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


def is_empty(value):
    return value is None or (
        isinstance(value, str) and value.strip() == ""
    )


def detect_type(cell):
    if cell.data_type == "f":
        return "formula"

    if cell.data_type == "e":
        return "excel_error"

    value = cell.value

    if isinstance(value, bool):
        return "boolean"

    if isinstance(value, (datetime, date, time)):
        return "date_or_time"

    if isinstance(value, (int, float)):
        return "number"

    if isinstance(value, str):
        return "text"

    return "other"


def inspect_excel(file_path):
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=False,
    )

    try:
        print(f"File: {file_path.name}")

        for sheet in workbook.worksheets:
            rows = sheet.iter_rows()
            header_cells = next(rows, ())

            columns = [
                {
                    "name": (
                        str(cell.value)
                        if not is_empty(cell.value)
                        else f"(unnamed column {index})"
                    ),
                    "missing": 0,
                    "types": Counter(),
                }
                for index, cell in enumerate(header_cells, start=1)
            ]

            data_rows = 0

            for row in rows:
                if all(is_empty(cell.value) for cell in row):
                    continue

                data_rows += 1

                for index, column in enumerate(columns):
                    if index >= len(row) or is_empty(row[index].value):
                        column["missing"] += 1
                    else:
                        value_type = detect_type(row[index])
                        column["types"][value_type] += 1

            print(f"\nSheet: {sheet.title}")
            print(f"Data rows: {data_rows}")

            for column in columns:
                types_summary = ", ".join(
                    f"{name}: {count}"
                    for name, count in sorted(column["types"].items())
                ) or "no nonempty values"

                missing = column["missing"]

                missing_percent = (
                    missing / data_rows * 100
                    if data_rows > 0
                    else 0
                )

                print(f"  Column: {column['name']}")
                print(
                    f"    Missing: {missing}/{data_rows} "
                    f"({missing_percent:.1f}%)"
                )
                print(f"    Types: {types_summary}")

    finally:
        workbook.close()


def main():
    parser = argparse.ArgumentParser(
        description="Profile columns in an Excel workbook."
    )
    parser.add_argument("file", type=Path)
    args = parser.parse_args()

    file_path = args.file.resolve()

    if not file_path.is_file():
        parser.error(f"File not found: {file_path}")

    if file_path.suffix.lower() != ".xlsx":
        parser.error("This version supports .xlsx files only.")

    inspect_excel(file_path)


if __name__ == "__main__":
    main()