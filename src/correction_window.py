import json
import os
import queue
import subprocess
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import messagebox, ttk
from tkinter.scrolledtext import ScrolledText
from uuid import uuid4

from openpyxl import load_workbook


PROJECT_FOLDER = Path(__file__).resolve().parent.parent
SCRIPT_FOLDER = Path(__file__).resolve().parent


class CorrectionWindow:
    def __init__(self, root, source, baseline):
        self.root = root
        self.source = Path(source).resolve()
        self.baseline = Path(baseline).resolve()
        self.events = queue.Queue()
        self.running = False
        self.preview_path = None

        workbook = load_workbook(self.source, read_only=True)

        try:
            self.headers = {}

            for sheet in workbook.worksheets:
                first_row = next(
                    sheet.iter_rows(
                        min_row=1,
                        max_row=1,
                        values_only=True,
                    ),
                    (),
                )

                self.headers[sheet.title] = [
                    value
                    for value in first_row
                    if isinstance(value, str) and value.strip()
                ]
        finally:
            workbook.close()

        root.title("FlowGuard AI — Reviewed Corrections")
        root.geometry("850x600")

        frame = ttk.Frame(root, padding=24)
        frame.pack(fill="both", expand=True)

        ttk.Label(
            frame,
            text="Reviewed Whitespace Correction",
            font=("Segoe UI", 18, "bold"),
        ).pack(anchor="w")

        ttk.Label(
            frame,
            text=(
                "Choose a column where removing outer spaces is valid. "
                "Output is a separate copy; source and approved stay unchanged."
            ),
            wraplength=780,
        ).pack(anchor="w", pady=12)

        self.sheet = tk.StringVar()
        self.column = tk.StringVar()

        ttk.Label(frame, text="Worksheet").pack(anchor="w")

        self.sheet_picker = ttk.Combobox(
            frame,
            textvariable=self.sheet,
            values=list(self.headers),
            state="readonly",
        )
        self.sheet_picker.pack(fill="x", pady=(4, 10))
        self.sheet_picker.bind(
            "<<ComboboxSelected>>",
            self.change_sheet,
        )

        ttk.Label(frame, text="Column").pack(anchor="w")

        self.column_picker = ttk.Combobox(
            frame,
            textvariable=self.column,
            state="readonly",
        )
        self.column_picker.pack(fill="x", pady=(4, 12))
        self.column_picker.bind(
            "<<ComboboxSelected>>",
            self.clear_preview,
        )

        actions = ttk.Frame(frame)
        actions.pack(anchor="w")

        self.preview_button = ttk.Button(
            actions,
            text="Preview",
            command=self.start_preview,
        )
        self.preview_button.pack(side="left", padx=(0, 10))

        self.apply_button = ttk.Button(
            actions,
            text="Create Corrected Copy",
            command=self.start_apply,
            state="disabled",
        )
        self.apply_button.pack(side="left")

        self.status = tk.StringVar(value="Choose worksheet and column")

        ttk.Label(
            frame,
            textvariable=self.status,
        ).pack(anchor="w", pady=12)

        self.log = ScrolledText(
            frame,
            font=("Consolas", 10),
            state="disabled",
        )
        self.log.pack(fill="both", expand=True)

        root.protocol("WM_DELETE_WINDOW", self.close)
        root.after(100, self.process_events)

    def clear_preview(self, event=None):
        self.preview_path = None
        self.apply_button.configure(state="disabled")

    def change_sheet(self, event=None):
        self.column.set("")
        self.column_picker.configure(
            values=self.headers.get(self.sheet.get(), [])
        )
        self.clear_preview()

    def write_log(self, text):
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def set_busy(self, busy):
        self.running = busy
        self.preview_button.configure(
            state="disabled" if busy else "normal"
        )
        self.sheet_picker.configure(
            state="disabled" if busy else "readonly"
        )
        self.column_picker.configure(
            state="disabled" if busy else "readonly"
        )
        self.apply_button.configure(state="disabled")

    def run_command(self, script, *arguments):
        environment = os.environ.copy()
        environment["PYTHONIOENCODING"] = "utf-8"

        result = subprocess.run(
            [
                sys.executable,
                str(SCRIPT_FOLDER / script),
                *[str(argument) for argument in arguments],
            ],
            cwd=PROJECT_FOLDER,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            env=environment,
            timeout=240,
            check=False,
        )

        self.events.put((
            "log",
            result.stdout + "\n" + result.stderr,
        ))

        return result.returncode

    def start_preview(self):
        if self.running:
            return

        if not self.sheet.get() or not self.column.get():
            messagebox.showwarning(
                "Choose column",
                "Choose a worksheet and column first.",
            )
            return

        self.clear_preview()
        self.set_busy(True)
        self.status.set("Creating preview…")

        threading.Thread(
            target=self.preview_job,
            args=(self.sheet.get(), self.column.get()),
            daemon=True,
        ).start()

    def preview_job(self, sheet, column):
        try:
            preview = PROJECT_FOLDER / "reports" / (
                f"trim_preview_{uuid4().hex}.json"
            )

            code = self.run_command(
                "preview_trim.py",
                self.source,
                "--sheet", sheet,
                "--column", column,
                "--output", preview,
            )

            if code != 0:
                raise RuntimeError("Preview failed. See log.")

            plan = json.loads(
                preview.read_text(encoding="utf-8")
            )

            self.events.put((
                "preview_done",
                (preview, len(plan["candidates"])),
            ))

        except Exception as error:
            self.events.put(("failed", str(error)))

    def start_apply(self):
        if self.running or self.preview_path is None:
            return

        confirmed = messagebox.askokcancel(
            "Confirm correction",
            "Remove outer spaces from the previewed text cells?\n\n"
            "A separate copy will be created and checked. "
            "Source and approved will not be replaced.",
        )

        if not confirmed:
            return

        preview = self.preview_path
        self.preview_path = None
        self.set_busy(True)
        self.status.set("Creating and checking corrected copy…")

        threading.Thread(
            target=self.apply_job,
            args=(preview,),
            daemon=True,
        ).start()

    def apply_job(self, preview):
        try:
            output = (
                PROJECT_FOLDER
                / "reports"
                / "corrected"
                / f"corrected_{uuid4().hex}.xlsx"
            )

            code = self.run_command(
                "apply_trim.py",
                self.source,
                preview,
                output,
                "--confirm",
            )

            if code != 0:
                raise RuntimeError("Correction failed. See log.")

            code = self.run_command(
                "check_workbook.py",
                output,
                self.baseline,
            )

            if code != 0:
                raise RuntimeError(
                    "Copy created but did not pass checks. "
                    f"Not approved. Location: {output}"
                )

            self.events.put(("apply_done", str(output)))

        except Exception as error:
            self.events.put(("failed", str(error)))

    def process_events(self):
        while True:
            try:
                event, value = self.events.get_nowait()
            except queue.Empty:
                break

            if event == "log":
                self.write_log(value)

            elif event == "preview_done":
                self.set_busy(False)
                preview, count = value
                self.status.set(f"Preview: {count} candidate cells")

                if count:
                    self.preview_path = preview
                    self.apply_button.configure(state="normal")

            elif event == "apply_done":
                self.set_busy(False)
                self.status.set("Corrected copy passed current checks")
                self.write_log(f"OUTPUT: {value}")
                self.write_log(
                    "Not published to approved. No refresh triggered."
                )

            elif event == "failed":
                self.set_busy(False)
                self.preview_path = None
                self.status.set("Stopped — review required")
                self.write_log(value)

        self.root.after(100, self.process_events)

    def close(self):
        if self.running:
            messagebox.showinfo(
                "Operation running",
                "Wait for this operation to finish before closing.",
            )
            return

        self.root.destroy()