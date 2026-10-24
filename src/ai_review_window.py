import json
import os
import queue
import subprocess
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText
from uuid import uuid4

from openpyxl import load_workbook

from inspect_excel import inspect_excel


PROJECT_FOLDER = Path(__file__).resolve().parent.parent
REVIEW_SCRIPT = Path(__file__).resolve().with_name(
    "review_profile.py"
)


class AIReviewWindow:
    def __init__(self, root, source):
        self.root = root
        self.source = source
        self.running = False
        self.events = queue.Queue()

        workbook = load_workbook(source, read_only=True)

        try:
            sheet_names = workbook.sheetnames
        finally:
            workbook.close()

        root.title("FlowGuard AI — Local AI Review")
        root.geometry("850x600")

        frame = ttk.Frame(root, padding=24)
        frame.pack(fill="both", expand=True)

        ttk.Label(
            frame,
            text="Local AI Review",
            font=("Segoe UI", 20, "bold"),
        ).pack(anchor="w")

        ttk.Label(
            frame,
            text=f"Source: {source.name}",
        ).pack(anchor="w", pady=(8, 16))

        ttk.Label(
            frame,
            text="Choose a worksheet:",
        ).pack(anchor="w")

        self.sheet = tk.StringVar(
            value=sheet_names[0] if sheet_names else ""
        )

        self.sheet_picker = ttk.Combobox(
            frame,
            textvariable=self.sheet,
            values=sheet_names,
            state="readonly",
        )
        self.sheet_picker.pack(fill="x", pady=(4, 12))

        ttk.Label(
            frame,
            text=(
                "Sends column names, types and missing-value counts "
                "to the local Ollama server. No cell values included.\n"
                "Metadata may still be sensitive. "
                "Suggestions are not approved rules."
            ),
            wraplength=780,
        ).pack(anchor="w", pady=(0, 12))

        self.button = ttk.Button(
            frame,
            text="Review Selected Sheet",
            command=self.start_review,
        )
        self.button.pack(anchor="w")

        self.status = tk.StringVar(value="Ready")

        ttk.Label(
            frame,
            textvariable=self.status,
        ).pack(anchor="w", pady=10)

        self.output = ScrolledText(
            frame,
            font=("Consolas", 10),
            state="disabled",
        )
        self.output.pack(fill="both", expand=True)

        root.protocol("WM_DELETE_WINDOW", self.close)
        root.after(100, self.process_events)

    def write_output(self, text):
        self.output.configure(state="normal")
        self.output.insert("end", text)
        self.output.see("end")
        self.output.configure(state="disabled")

    def start_review(self):
        if self.running or not self.sheet.get():
            return

        consent = messagebox.askokcancel(
            "Local metadata review",
            "Review this sheet's metadata using local Ollama?\n\n"
            "Keep Ollama in local-only mode. "
            "No Excel edits will be made.",
        )

        if not consent:
            return

        self.running = True
        self.button.configure(state="disabled")
        self.sheet_picker.configure(state="disabled")
        self.status.set("Profiling and reviewing…")
        self.write_output("\n--- LOCAL AI REVIEW ---\n")

        threading.Thread(
            target=self.run_review,
            args=(self.sheet.get(),),
            daemon=True,
        ).start()

    def run_review(self, sheet_name):
        try:
            # نقرأ المصدر من جديد، بدل استخدام وصف قديم
            profile = inspect_excel(self.source)

            report_folder = PROJECT_FOLDER / "reports"
            report_folder.mkdir(parents=True, exist_ok=True)

            profile_path = report_folder / (
                f"ai_profile_{uuid4().hex}.json"
            )

            profile_path.write_text(
                json.dumps(
                    profile,
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )

            environment = os.environ.copy()
            environment["PYTHONIOENCODING"] = "utf-8"

            try:
                result = subprocess.run(
                    [
                        sys.executable,
                        str(REVIEW_SCRIPT),
                        str(profile_path),
                        "--sheet",
                        sheet_name,
                    ],
                    cwd=PROJECT_FOLDER,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    env=environment,
                    timeout=240,
                    check=False,
                )

                text = result.stdout

                if result.stderr:
                    text += "\n" + result.stderr

                self.events.put(("output", text))
                self.events.put((
                    "done",
                    "Review finished — inspect output; "
                    "no rules approved.",
                ))

            finally:
                profile_path.unlink(missing_ok=True)

        except Exception as error:
            self.events.put((
                "output",
                f"Review could not complete: {error}\n",
            ))
            self.events.put(("done", "Review failed — see details"))

    def process_events(self):
        while True:
            try:
                event, value = self.events.get_nowait()
            except queue.Empty:
                break

            if event == "output":
                self.write_output(value)

            elif event == "done":
                self.running = False
                self.status.set(value)
                self.button.configure(state="normal")
                self.sheet_picker.configure(state="readonly")

        self.root.after(100, self.process_events)

    def close(self):
        if self.running:
            messagebox.showinfo(
                "Review running",
                "Wait for the review to finish before closing.",
            )
            return

        self.root.destroy()


def main():
    root = tk.Tk()
    root.withdraw()

    source = filedialog.askopenfilename(
        title="Choose Excel for local AI review",
        initialdir=PROJECT_FOLDER / "data" / "incoming",
        filetypes=[("Excel workbook", "*.xlsx")],
    )

    if not source:
        root.destroy()
        return

    try:
        AIReviewWindow(root, Path(source))
        root.deiconify()

    except Exception as error:
        messagebox.showerror("Could not open workbook", str(error))
        root.destroy()
        return

    root.mainloop()


if __name__ == "__main__":
    main()