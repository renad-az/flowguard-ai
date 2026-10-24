import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from approve_workbook import file_hash


def main():
    parser = argparse.ArgumentParser(
        description="Preview outer-space removal in selected columns."
    )

    parser.add_argument("source", type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument(
        "--column",
        action="append",
        required=True,
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("reports/trim_preview.json"),
    )

    args = parser.parse_args()

    source = args.source.resolve()
    output = args.output.resolve()

    if source.suffix.lower() != ".xlsx":
        parser.error("Only .xlsx files are supported.")

    if not source.is_file():
        parser.error("Source file not found.")

    if output == source:
        parser.error("Output must not overwrite the source.")

    try:
        before_hash = file_hash(source)
        workbook = load_workbook(source, read_only=True)

        candidates = []

        try:
            if args.sheet not in workbook.sheetnames:
                raise ValueError("Selected sheet not found.")

            sheet = workbook[args.sheet]
            rows = sheet.iter_rows()
            headers = [
                cell.value
                for cell in next(rows, ())
            ]

            selected = set(args.column)
            positions = {}

            for column in selected:
                matches = [
                    index
                    for index, name in enumerate(headers)
                    if name == column
                ]

                if len(matches) != 1:
                    raise ValueError(
                        f"Column must exist exactly once: {column}"
                    )

                positions[column] = matches[0]

            for row in rows:
                for column, index in positions.items():
                    if index >= len(row):
                        continue

                    cell = row[index]

                    # نفحص النصوص فقط؛ لا نغيّر الصيغ أو الأرقام
                    if (
                        cell.data_type != "s"
                        or not isinstance(cell.value, str)
                    ):
                        continue

                    trimmed = cell.value.strip()

                    # النص المكوّن من مسافات فقط يحتاج قرارًا منفصلًا
                    if not trimmed:
                        continue

                    if trimmed != cell.value:
                        candidates.append({
                            "sheet": args.sheet,
                            "column": column,
                            "cell": cell.coordinate,
                            "action": "trim_outer_spaces",
                        })

        finally:
            workbook.close()

        if file_hash(source) != before_hash:
            raise ValueError(
                "Source changed during inspection. Run again."
            )

        preview = {
            "version": 1,
            "source_sha256": before_hash,
            "approved_for_execution": False,
            "candidates": candidates,
        }

        output.parent.mkdir(parents=True, exist_ok=True)

        output.write_text(
            json.dumps(
                preview,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        print("Candidate cells:", len(candidates))

        for candidate in candidates[:20]:
            print(
                "REVIEW:",
                candidate["sheet"],
                candidate["cell"],
                candidate["column"],
            )

        if len(candidates) > 20:
            print("Additional locations are in the preview file.")

        print("Preview saved:", output)
        print("No cell values included in the preview.")
        print("No Excel files modified. No AI request sent.")

        return 0

    except Exception as error:
        print("PREVIEW FAILED:", type(error).__name__, str(error))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())