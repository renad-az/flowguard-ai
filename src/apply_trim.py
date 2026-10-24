import argparse
import hashlib
import json
import os
import tempfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import load_workbook


def file_hash(path):
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def main():
    parser = argparse.ArgumentParser(
        description="Create a verified corrected copy and audit record."
    )

    parser.add_argument("source", type=Path)
    parser.add_argument("preview", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--confirm", action="store_true")
    args = parser.parse_args()

    if not args.confirm:
        parser.error("Review the preview, then use --confirm.")

    source = args.source.resolve()
    preview_path = args.preview.resolve()
    output = args.output.resolve()
    audit_path = output.with_suffix(".audit.json")

    if output in (source, preview_path):
        parser.error("Output must be a separate file.")

    if output.suffix.lower() != ".xlsx":
        parser.error("Output must use .xlsx.")

    if output.exists() or audit_path.exists():
        parser.error("Output or audit already exists. Choose a new name.")

    workbook = None
    temporary = None
    audit_temporary = None
    output_published = False
    audit_published = False

    try:
        plan = json.loads(
            preview_path.read_text(encoding="utf-8")
        )

        if plan.get("version") != 1:
            raise ValueError("Unsupported preview version.")

        source_bytes = source.read_bytes()
        source_hash = hashlib.sha256(source_bytes).hexdigest()

        if source_hash != plan["source_sha256"]:
            raise ValueError(
                "Source changed since preview. Create a new preview."
            )

        candidates = plan["candidates"]

        if not candidates:
            raise ValueError("No corrections in this preview.")

        workbook = load_workbook(
            BytesIO(source_bytes),
            data_only=False,
        )

        original_sheets = list(workbook.sheetnames)
        expected_values = {}
        actions = []
        seen = set()

        for candidate in candidates:
            if candidate["action"] != "trim_outer_spaces":
                raise ValueError("Unsupported correction.")

            sheet_name = candidate["sheet"]
            coordinate = candidate["cell"]
            column_name = candidate["column"]

            if sheet_name not in workbook.sheetnames:
                raise ValueError("Preview sheet not found.")

            sheet = workbook[sheet_name]
            cell = sheet[coordinate]

            if cell.row < 2:
                raise ValueError("Header corrections are not allowed.")

            if sheet.cell(1, cell.column).value != column_name:
                raise ValueError("Preview column does not match.")

            key = (sheet_name, cell.coordinate)

            if key in seen:
                raise ValueError("Duplicate correction.")

            seen.add(key)

            if (
                cell.data_type != "s"
                or not isinstance(cell.value, str)
            ):
                raise ValueError("Only text cells can be trimmed.")

            trimmed = cell.value.strip()

            if not trimmed or trimmed == cell.value:
                raise ValueError("Candidate is not a valid trim.")

            cell.value = trimmed
            expected_values[key] = trimmed

            actions.append({
                "sheet": sheet_name,
                "column": column_name,
                "cell": cell.coordinate,
                "action": "trim_outer_spaces",
            })

        output.parent.mkdir(parents=True, exist_ok=True)

        with tempfile.NamedTemporaryFile(
            dir=output.parent,
            prefix=".trim-pending-",
            suffix=".xlsx",
            delete=False,
        ) as file:
            temporary = Path(file.name)

        workbook.save(str(temporary))
        workbook.close()
        workbook = None

        with ZipFile(temporary, "r") as archive:
            if archive.testzip() is not None:
                raise ValueError("Workbook archive integrity failed.")

        verified = load_workbook(
            temporary,
            read_only=True,
            data_only=False,
        )

        try:
            if verified.sheetnames != original_sheets:
                raise ValueError("Worksheet list changed unexpectedly.")

            for (sheet_name, coordinate), expected in expected_values.items():
                if verified[sheet_name][coordinate].value != expected:
                    raise ValueError(
                        f"Correction verification failed: "
                        f"{sheet_name}!{coordinate}"
                    )
        finally:
            verified.close()

        if file_hash(source) != source_hash:
            raise ValueError(
                "Source changed during correction. Preview again."
            )

        audit = {
            "version": 1,
            "operation_id": uuid4().hex,
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "confirmation": "explicit_confirm_flag",
            "source_sha256": source_hash,
            "output_sha256": file_hash(temporary),
            "corrected_cells": len(actions),
            "actions": actions,
            "checks": {
                "archive_integrity": "passed",
                "worksheet_names": "passed",
                "corrected_cell_values": "passed",
                "baseline_validation": "not_run_by_this_script",
            },
            "scope": (
                "Separate corrected copy only; "
                "not approved and no refresh triggered."
            ),
        }

        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=output.parent,
            prefix=".audit-pending-",
            suffix=".json",
            delete=False,
        ) as file:
            audit_temporary = Path(file.name)
            json.dump(audit, file, ensure_ascii=False, indent=2)

        # نشر بدون استبدال ملفات موجودة
        os.link(temporary, output)
        output_published = True

        os.link(audit_temporary, audit_path)
        audit_published = True

        print("CORRECTED COPY:", output)
        print("AUDIT RECORD:", audit_path)
        print("Cells trimmed:", len(actions))
        print("Archive and corrected-cell checks passed.")
        print("Baseline validation must be run separately.")
        print("No cell values stored in the audit record.")
        print("Original source and approved workbook unchanged.")
        print("No Power BI refresh triggered.")
        return 0

    except Exception as error:
        print("CORRECTION FAILED:", type(error).__name__, str(error))

        if output_published and not audit_published:
            print(
                "A corrected copy was created, but audit publication "
                "failed. Review manually; do not approve automatically."
            )

        return 1

    finally:
        if workbook is not None:
            workbook.close()

        if temporary is not None:
            temporary.unlink(missing_ok=True)

        if audit_temporary is not None:
            audit_temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    raise SystemExit(main())